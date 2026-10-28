"""Read the two research workbooks without modifying them; export static web data.

Requires openpyxl. Run from any directory. See README.md for update instructions.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

SITE = Path(__file__).resolve().parents[1]
ALIASES = {"BCIC2A": "BCIC-2A", "BCIC2B": "BCIC-2B", "SEEDIV": "SEED-IV", "PBCIW": "pBCIW"}
DATASETS = [
    {"id": "BCIC-2A", "label": "BCIC-2A", "task": "Motor imagery", "shortTask": "MI", "classes": 4},
    {"id": "BCIC-2B", "label": "BCIC-2B", "task": "Motor imagery", "shortTask": "MI", "classes": 2},
    {"id": "SEED-IV", "label": "SEED-IV", "task": "Emotion recognition", "shortTask": "Emotion", "classes": 4},
    {"id": "pBCIW", "label": "pBCIW", "task": "Cognitive workload", "shortTask": "Workload", "classes": None},
]
SETTINGS = ("frozen", "finetuned", "scratch")
NUMBER = r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?"


def canonical(value):
    key = re.sub(r"[^A-Za-z0-9]", "", str(value)).upper()
    if key not in ALIASES:
        raise ValueError(f"Unknown dataset: {value}")
    return ALIASES[key]


def numeric(value, location):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"Expected a finite numeric value at {location}; got {value!r}")
    return float(value)


def measurement(value, location, required=True):
    if value is None:
        if required:
            raise ValueError(f"Missing required measurement: {location}")
        return None
    if isinstance(value, (int, float)):
        return {"mean": numeric(value, location), "spread": None}
    match = re.fullmatch(rf"\s*({NUMBER})(?:\s*(?:±|\+/-)\s*({NUMBER}))?\s*", str(value))
    if not match:
        raise ValueError(f"Unrecognized measurement at {location}: {value!r}")
    result = {"mean": float(match[1]), "spread": float(match[2]) if match[2] else None}
    if not all(math.isfinite(v) for v in result.values() if v is not None):
        raise ValueError(f"Non-finite measurement at {location}")
    if result["spread"] is not None and result["spread"] < 0:
        raise ValueError(f"Negative spread at {location}")
    return result


def derive_metrics(kf, kt, kr):
    if kt <= 0:
        return {"prr": None, "acs": None, "pvs": None}
    prr, acs = kf / kt, kr / kt
    pvs = None if prr < 0 else (0.0 if acs >= 1 else 2 * prr * (1 - acs) / (prr + 1 - acs))
    return {"prr": prr, "acs": acs, "pvs": pvs}


def export(performance_path, metrics_path, output):
    # data_only=True reads cached results, never executes formulas or external links.
    metrics_book = load_workbook(metrics_path, data_only=True, read_only=True)
    metric_records = {}
    provenance = {}
    for sheet in metrics_book:
        dataset = None
        for cells in sheet.iter_rows():
            values = [c.value for c in cells]
            if not values or values[0] is None or values[0] == "Method":
                continue
            if all(v is None for v in values[1:]):
                dataset = canonical(values[0])
                continue
            if dataset is None:
                raise ValueError(f"A dataset heading must precede {sheet.title}!A{cells[0].row}")
            model = str(values[0]).strip()
            key = (dataset, model)
            if key in metric_records:
                raise ValueError(f"Duplicate metrics: {key}")
            metric_records[key] = {m: numeric(values[i], f"{sheet.title}!{cells[i].coordinate}")
                                   for i, m in enumerate(("prr", "acs", "pvs"), start=1)}
            provenance[key] = {"file": metrics_path.name, "sheet": sheet.title,
                               "range": f"B{cells[0].row}:D{cells[0].row}"}
    metrics_book.close()

    performance_book = load_workbook(performance_path, data_only=True, read_only=True)
    records, missing, seen = [], [], set()
    for sheet in performance_book:
        dataset = canonical(sheet.title)
        aux_header = str(sheet.cell(1, 8).value)
        aux_key = "auc" if "AUC" in aux_header.upper() else "f1"
        for row in sheet.iter_rows(min_row=3):
            if row[0].value in (None, "Method"):
                break
            model = str(row[0].value).strip()
            key = (dataset, model)
            if key in seen:
                raise ValueError(f"Duplicate performance: {key}")
            seen.add(key)
            if key not in metric_records:
                raise ValueError(f"Metrics not supplied for {key}")
            performance = {}
            for metric, offset in (("bac", 1), ("kappa", 4), (aux_key, 7)):
                performance[metric] = {}
                for j, setting in enumerate(SETTINGS):
                    cell = row[offset + j]
                    loc = f"{performance_path.name}/{sheet.title}!{cell.coordinate}"
                    val = measurement(cell.value, loc, required=metric in ("bac", "kappa"))
                    if val is None:
                        missing.append({"dataset": dataset, "model": model, "metric": metric,
                                        "setting": setting, "source": loc})
                    performance[metric][setting] = val
            k = performance["kappa"]
            calculated = derive_metrics(*(k[s]["mean"] for s in SETTINGS))
            supplied = metric_records[key]
            for m, val in calculated.items():
                if val is None or not math.isclose(val, supplied[m], rel_tol=1e-9, abs_tol=1e-10):
                    raise ValueError(f"Metric mismatch for {key} {m}: supplied={supplied[m]}, derived={val}")
            records.append({"dataset": dataset, "model": model, "metrics": supplied,
                            "performance": performance,
                            "source": {"metrics": provenance[key],
                                       "performance": {"file": performance_path.name,
                                                       "sheet": sheet.title,
                                                       "range": f"B{row[0].row}:J{row[0].row}"}}})
    performance_book.close()
    if seen != set(metric_records):
        raise ValueError(f"Unmatched metric rows: {set(metric_records) - seen}")
    if not records:
        raise ValueError("No results imported")
    model_ids = list(dict.fromkeys(r["model"] for r in records))
    dataset_ids = {d["id"] for d in DATASETS}
    if seen != {(d, m) for d in dataset_ids for m in model_ids}:
        raise ValueError("Expected a complete model-by-dataset result grid")
    files = [{"name": p.name, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}
             for p in (performance_path, metrics_path)]
    payload = {"schemaVersion": 1, "importedAt": datetime.now(timezone.utc).isoformat(),
               "protocol": "LOSO", "aggregation": "Ratios of reported mean Cohen's kappa scores",
               "spreadLabel": "Reported ± value; dispersion definition not specified in source workbooks",
               "datasets": DATASETS, "modelOrder": model_ids, "sources": files,
               "missingMeasurements": missing, "records": records}
    output.mkdir(parents=True, exist_ok=True)
    encoded = json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False)
    (output / "results.json").write_text(encoded + "\n", encoding="utf-8")
    # Also support a double-click file:// preview with no fetch/CORS requirements.
    (output / "results.js").write_text("window.ATLAS_DATA = " + encoded.replace("</", "<\\/") + ";\n", encoding="utf-8")
    out = io.StringIO(newline="")
    fields = ["dataset", "model", "PRR", "ACS", "PVS"]
    for metric in ("kappa", "bac", "f1", "auc"):
        for setting in SETTINGS:
            fields.extend((f"{metric}_{setting}", f"{metric}_{setting}_spread"))
    writer = csv.DictWriter(out, fieldnames=fields)
    writer.writeheader()
    for record in records:
        flat = {"dataset": record["dataset"], "model": record["model"],
                **{m.upper(): v for m, v in record["metrics"].items()}}
        for metric, settings in record["performance"].items():
            for setting, value in settings.items():
                flat[f"{metric}_{setting}"] = value["mean"] if value else ""
                flat[f"{metric}_{setting}_spread"] = value["spread"] if value else ""
        writer.writerow(flat)
    (output / "results.csv").write_text(out.getvalue(), encoding="utf-8-sig")
    print(f"Imported {len(records)} model/dataset records; validated {len(records) * 3} proposed metric values.")
    print(f"Missing optional measurements: {len(missing)}. Original workbooks were not modified.")
    return payload


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--performance", type=Path, default=SITE.parent / "logs" / "结果汇总.xlsx")
    parser.add_argument("--metrics", type=Path, default=SITE.parent / "logs" / "新指标单独汇总.xlsx")
    parser.add_argument("--output", type=Path, default=SITE / "data")
    args = parser.parse_args()
    export(args.performance.resolve(), args.metrics.resolve(), args.output.resolve())
